# -*- coding: utf-8 -*-
"""
Importa el Excel V3 y genera los JSON de datos de la plataforma.

Salida (uno por dominio, editables desde la UI):
    datos/equipos.json       catalogo de equipos (9)
    datos/params.json        parametros de entrada agrupados (editables)
    datos/secciones.json     secciones de labor de la mina
    datos/referencias.json   fuentes con pagina

Uso:
    python -m engine.importar [ruta_al_xlsx]
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

from .modelo import DATOS, RAIZ            # misma carpeta de trabajo que el motor

XLSX_DEFECTO = RAIZ.parent / "Caso_Practico_Scoop_Uchucchacua v3.xlsx"

# ---------------------------------------------------------------- utilidades

def _limpiar(v: Any) -> Any:
    if isinstance(v, str):
        return v.strip()
    return v


def _slug(texto: str, limite: int = 52) -> str:
    """'Precio diesel (US$/gal)' -> 'precio_diesel'"""
    t = texto.lower()
    t = re.sub(r"\(.*?\)", "", t)                       # fuera parentesis
    t = (t.replace("á", "a").replace("é", "e").replace("í", "i")
           .replace("ó", "o").replace("ú", "u").replace("ñ", "n"))
    t = re.sub(r"[^a-z0-9]+", "_", t).strip("_")
    return re.sub(r"_+", "_", t)[:limite].strip("_")


def _slug_unico(etiqueta: str, usados: set[str]) -> str:
    """Slug garantizado unico dentro del documento.

    'Produccion objetivo (TCS/dia)' y 'Produccion objetivo (TCS/guardia)'
    comparten raiz; el segundo toma el sufijo de su unidad.
    """
    base = _slug(etiqueta)
    if base not in usados:
        usados.add(base)
        return base
    unidad = _unidad(etiqueta)
    if unidad:
        cand = f"{base}_{_slug(unidad, 18)}"
        if cand and cand not in usados:
            usados.add(cand)
            return cand
    i = 2
    while f"{base}_{i}" in usados:
        i += 1
    usados.add(f"{base}_{i}")
    return f"{base}_{i}"


def _unidad(label: str) -> str:
    m = re.search(r"\(([^)]*)\)\s*$", label)
    return m.group(1) if m else ""


def _tipo(label: str, fmt: str, valor: Any) -> str:
    """porcentaje | moneda | entero | decimal | texto"""
    if isinstance(valor, str):
        return "texto"
    if "%" in (fmt or ""):
        return "porcentaje"
    if "$" in (fmt or "") or "US$" in label:
        return "moneda"
    if isinstance(valor, int) or (isinstance(valor, float) and valor.is_integer() and abs(valor) >= 1):
        u = _unidad(label)
        if u in ("mm", "m", "") or "dias" in label.lower() or "horas" in label.lower():
            return "entero"
    return "decimal"


def _rango(tipo: str, valor: Any) -> tuple[float | None, float | None, float | None]:
    """(min, max, paso) sugeridos para el editor."""
    if not isinstance(valor, (int, float)):
        return None, None, None
    if tipo == "porcentaje":
        return 0.0, 1.0, 0.005
    if valor == 0:
        return 0.0, 1.0, 0.01
    lo, hi = 0.0, round(abs(valor) * 3, 6)
    paso = 10 ** (len(str(int(abs(valor)))) - 3) if abs(valor) >= 100 else 0.01
    return lo, hi, max(paso, 0.001)


# ---------------------------------------------------------------- equipos

CAMPOS_EQUIPO = {
    1: ("n", "int"), 2: ("modelo", "str"), 3: ("fabricante", "str"),
    4: ("energia", "str"), 5: ("situacion", "str"),
    6: ("cuchara_m3", "float"), 7: ("cuchara_yd3", "float"),
    8: ("carga_util_kg", "float"), 9: ("ancho_mm", "float"),
    10: ("alto_mm", "float"), 11: ("largo_mm", "float"),
    12: ("potencia_kw", "float"), 13: ("potencia_hp", "float"),
    14: ("peso_kg", "float"), 15: ("fuerza_arranque_kg", "float"),
    16: ("radio_giro_mm", "float"), 17: ("vel_max_kmh", "float"),
    18: ("autonomia", "str"), 19: ("precio_usd", "float"),
    20: ("vida_util_h", "float"), 21: ("consumo", "float"),
    22: ("flexibilidad", "float"), 23: ("fuente", "str"),
}

# clave de render 3D segun tipo de energia
RENDER = {
    "Diesel":          {"clase": "diesel",   "acento": "warn", "escape": True,  "cable_m": 0,   "bateria": False},
    "Electrico-cable": {"clase": "cable",    "acento": "accent", "escape": False, "cable_m": 120, "bateria": False},
    "Bateria":         {"clase": "bateria",  "acento": "ok",   "escape": False, "cable_m": 0,   "bateria": True},
}


def leer_equipos(wv) -> list[dict]:
    ws = wv["Equipos"]
    equipos = []
    r = 5
    while ws.cell(r, 1).value is not None:
        eq: dict[str, Any] = {}
        for col, (clave, tipo) in CAMPOS_EQUIPO.items():
            v = _limpiar(ws.cell(r, col).value)
            if v is None:
                eq[clave] = None
            elif tipo == "int":
                eq[clave] = int(v)
            elif tipo == "float":
                eq[clave] = float(v) if isinstance(v, (int, float)) else None
            else:
                eq[clave] = str(v)
        eq["id"] = f"E{eq['n']:02d}"
        eq["render"] = RENDER.get(eq.get("energia") or "", RENDER["Diesel"])
        # unidad del consumo: diesel en gal/h, electricos en kWh/h
        eq["consumo_unidad"] = "gal/h" if eq.get("energia") == "Diesel" else "kWh/h"
        equipos.append(eq)
        r += 1
    return equipos


# ---------------------------------------------------------------- parametros

def leer_params(wf, wv) -> list[dict]:
    """Devuelve los grupos con sus campos editables (las celdas amarillas)."""
    ws, wsv = wf["Datos_Entrada"], wv["Datos_Entrada"]
    grupos: list[dict] = []
    actual: dict | None = None
    usados: set[str] = set()

    for r in range(1, ws.max_row + 1):
        etiqueta = _limpiar(ws.cell(r, 1).value)
        cruda = ws.cell(r, 2).value
        valor = wsv.cell(r, 2).value
        nota = _limpiar(ws.cell(r, 4).value)

        if etiqueta is None:
            continue

        es_formula = isinstance(cruda, str) and cruda.startswith("=")
        es_cabecera = cruda is None and valor is None

        if es_cabecera:
            if r <= 2:                       # titulo y subtitulo de la hoja
                continue
            actual = {"grupo": etiqueta, "slug": _slug(etiqueta), "campos": []}
            grupos.append(actual)
            continue

        if actual is None:
            continue

        fmt = ws.cell(r, 2).number_format or ""
        clave = _slug_unico(etiqueta, usados)
        if es_formula:
            actual["campos"].append({
                "clave": clave, "label": etiqueta, "celda": f"B{r}",
                "editable": False, "derivado": True,
                "valor": valor, "unidad": _unidad(etiqueta),
                "tipo": _tipo(etiqueta, fmt, valor), "nota": nota,
            })
            continue

        tipo = _tipo(etiqueta, fmt, valor)
        lo, hi, paso = _rango(tipo, valor)
        actual["campos"].append({
            "clave": clave, "label": etiqueta, "celda": f"B{r}",
            "editable": True, "derivado": False,
            "valor": valor, "unidad": _unidad(etiqueta), "tipo": tipo,
            "min": lo, "max": hi, "paso": paso, "nota": nota,
        })

    return [g for g in grupos if g["campos"]]


def aplanar(grupos: list[dict]) -> dict[str, Any]:
    """{clave: valor} de todos los campos editables, para el motor."""
    plano = {}
    for g in grupos:
        for c in g["campos"]:
            if c["editable"]:
                plano[c["clave"]] = c["valor"]
    return plano


# ---------------------------------------------------------------- secciones

def leer_secciones(wv) -> list[dict]:
    ws = wv["Restriccion_Seccion"]
    secciones = []
    for r in range(6, 10):
        nombre = _limpiar(ws.cell(r, 1).value)
        ancho, alto = ws.cell(r, 3).value, ws.cell(r, 4).value
        if nombre and isinstance(ancho, (int, float)):
            secciones.append({
                "clave": _slug(nombre), "nombre": nombre,
                "ancho_mm": float(ancho), "alto_mm": float(alto),
                "critica": "produccion" in nombre.lower(),
            })
    return secciones


# ---------------------------------------------------------------- referencias

def leer_referencias(wv) -> list[dict]:
    ws = wv["Referencias_Costos"]
    refs = []
    for r in range(5, ws.max_row + 1):
        concepto = _limpiar(ws.cell(r, 1).value)
        if not concepto or concepto.lower().startswith(("empresa", "contactos")):
            continue
        valor = _limpiar(ws.cell(r, 2).value)
        if valor is None:
            continue
        refs.append({
            "concepto": concepto, "valor": str(valor),
            "fuente": str(_limpiar(ws.cell(r, 3).value) or ""),
            "tabla": str(_limpiar(ws.cell(r, 4).value) or ""),
            "pagina": str(_limpiar(ws.cell(r, 5).value) or ""),
            "cita": str(_limpiar(ws.cell(r, 6).value) or ""),
            "verificar": str(_limpiar(ws.cell(r, 7).value) or ""),
        })
    return refs


# ---------------------------------------------------------------- escritura

def _guardar(nombre: str, obj: Any) -> Path:
    DATOS.mkdir(parents=True, exist_ok=True)
    ruta = DATOS / nombre
    ruta.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    return ruta


CAMPOS_FUERA_DEL_EXCEL = ("radio_giro_int_mm", "radio_giro_int_estimado")


def _conservar_radio_interior(equipos: list[dict]) -> None:
    """El Excel solo trae el radio exterior: el interior se conserva del JSON."""
    ruta = DATOS / "equipos.json"
    if not ruta.exists():
        return
    previos = {e["n"]: e for e in json.loads(ruta.read_text(encoding="utf-8"))["equipos"]}
    for eq in equipos:
        for campo in CAMPOS_FUERA_DEL_EXCEL:
            if campo in previos.get(eq["n"], {}):
                eq[campo] = previos[eq["n"]][campo]


def importar(xlsx: Path | None = None) -> dict[str, Path]:
    xlsx = Path(xlsx) if xlsx else XLSX_DEFECTO
    if not xlsx.exists():
        raise FileNotFoundError(f"No se encontro el Excel: {xlsx}")

    wf = openpyxl.load_workbook(xlsx, data_only=False)
    wv = openpyxl.load_workbook(xlsx, data_only=True)

    equipos = leer_equipos(wv)
    _conservar_radio_interior(equipos)
    grupos = leer_params(wf, wv)
    secciones = leer_secciones(wv)
    referencias = leer_referencias(wv)

    salidas = {
        "equipos": _guardar("equipos.json", {
            "origen": xlsx.name, "n": len(equipos), "equipos": equipos}),
        "params": _guardar("params.json", {
            "origen": xlsx.name, "grupos": grupos, "valores": aplanar(grupos)}),
        "secciones": _guardar("secciones.json", {
            "origen": xlsx.name, "secciones": secciones}),
        "referencias": _guardar("referencias.json", {
            "origen": xlsx.name, "n": len(referencias), "referencias": referencias}),
    }
    return salidas


if __name__ == "__main__":
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    out = importar(ruta)
    for k, v in out.items():
        print(f"  {k:12} -> {v.relative_to(RAIZ)}")
